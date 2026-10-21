# -*- coding: utf-8 -*-
"""
Готовое Python-решение для выделения и классификации мероприятий
благотворительного фонда «Дети Марии» на основе выгрузки постов ВКонтакте.

Функционал скрипта:
1. Загружает выгрузку постов (Excel .xlsx) и очищает тексты от пустых записей и служебной разметки ВК.
2. Отфильтровывает только посты о реальных мероприятиях (анонсы, приглашения, репортажи о занятиях,
   ярмарках, лагерях, выставках, спектаклях, праздниках и спортивных событиях), исключая некрологи,
   общие поздравления с календарными датами, грантовые релизы и сборы без событий.
3. Определяет основной и дополнительные типы каждого мероприятия по взвешенной модели регулярных выражений.
4. Выделяет целевую аудиторию (воспитанники ЦССВ/ДДИ, жители ПНИ/социальных домов и люди с ОВЗ,
   выпускники интернатов, семьи и дети Беслана, широкая аудитория, волонтёры) и конкретные учреждения.
5. Определяет востребованные навыки и компетенции волонтёров для каждого мероприятия.
6. Формирует многостраничный аналитический Excel-отчёт с KPI, сводными таблицами, диаграммами
   и полной размеченной базой мероприятий с кликабельными ссылками на посты.

Результат — «materials/02-фонд/вк-группа/Анализ мероприятий ВК 2011–2026.xlsx»,
исходные данные для будущего календаря мероприятий координатора.

Запуск из корня репозитория (по умолчанию читает mariaschildren_posts.xlsx и
перезаписывает таблицу анализа):
    python3 scripts/analyze_mariaschildren_events.py
    python3 scripts/analyze_mariaschildren_events.py --input <посты.xlsx> --output <отчёт.xlsx>
Нужны pandas и openpyxl.
"""

import argparse
import pathlib
import re
from collections import Counter

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent

# ---------------------------------------------------------------------------
# 1. СЛОВАРИ И ПРАВИЛА КЛАССИФИКАЦИИ (НА ОСНОВЕ ЛЕКСИКИ БФ «ДЕТИ МАРИИ»)
# ---------------------------------------------------------------------------

# Категории мероприятий: (Название, Высоковесные маркеры [вес 3], Средневесные маркеры [вес 1], Краткое описание формата)
EVENT_CATEGORIES = [
    (
        "Ярмарки, фестивали и благотворительные базары",
        r"(?<![а-яё])(ярмарк\w*|базар\w*|sherst_bazar|шерстян\w+\s+базар\w*|душевн\w+\s+bazar|душевн\w+\s+базар\w*|маркет\w*|аукцион\w*|благотворительн\w+\s+распродаж\w*|фестивал\w*)(?![а-яё])",
        r"(?<![а-яё])(стенд\w*\s+фонда|сувенир\w*\s+за\s+пожертвовани\w*)(?![а-яё])",
        "Благотворительные маркеты («Шерстяной Базар», «Душевный Bazar», «Жить красиво»), аукционы, фестивали"
    ),
    (
        "Выездные лагеря, дача и сезонные программы",
        r"(?<![а-яё])(лагер\w*|сорочан\w*|переделкин\w*|загородн\w+\s+мастерск\w*|на\s+дач[еуи]|дачн\w+\s+выезд\w*|летн\w+\s+арт-программ\w*|летн\w+\s+программ\w*|выездн\w+\s+программ\w*|поход\w*\s+на\s+байдарк\w*|байдарк\w*|остров\w*\s+в\s+сорочан\w*|арт-резиденци\w*|поездк\w+\s+в\s+беслан\w*|программ\w+\s+в\s+беслан\w*)(?![а-яё])",
        r"(?<![а-яё])(заезд\w*|смен[аеыу]\s+в|палатк\w*|у\s+костра)(?![а-яё])",
        "Летние арт-резиденции, лагеря в Сорочанах и Беслане, субботние выезды в загородную мастерскую (Переделкино)"
    ),
    (
        "Выставки, экскурсии и культурные выходы",
        r"(?<![а-яё])(выставк\w*|экспозици\w*|вернисаж\w*|экскурси\w*|медиаторск\w+\s+тур\w*|в\s+музе[ейи]|поход\w*\s+в\s+музей|посети\w+\s+музей|третьяков\w*|пушкинск\w+\s+музей|гэс-2|гэс–2|планетари\w*|зоопарк\w*|океанариум\w*|ботаническ\w+\s+сад\w*)(?![а-яё])",
        r"(?<![а-яё])(прогулк\w+\s+по|выход\w*\s+в\s+город|галере\w*)(?![а-яё])",
        "Художественные выставки работ подопечных (ГЭС-2, «Тверская 15»), медиаторские туры, экскурсии в музеи"
    ),
    (
        "Кулинарные, образовательные и социально-бытовые занятия",
        r"(?<![а-яё])(кулинар\w+|готови\w+\s+вместе|учимся\s+готовить|пекли|испекли|выпечк\w+|английск\w+\s+язык\w*|компьютерн\w+\s+курс\w*|компьютерн\w+\s+заняти\w*|подготовк\w+\s+к\s+школе|правил\w+\s+безопасност\w*|безопасность\s+без|подозрительн\w+\s+предмет\w*|бесед\w+\s+о\s+безопасност\w*|киноклуб\w*|книжн\w+\s+клуб\w*|финансов\w+\s+грамотност\w*|заняти\w+\s+по\s+кулинар\w*|ориентировани\w+\s+в\s+городе|хочу\s+больше\s+знать|настольн\w+\s+игр\w*|элиас|викторин\w*|квест\w*)(?![а-яё])",
        r"(?<![а-яё])(рецепт\w+|урок\w*|лекци\w*|семинар\w*|смотрим\s+фильм|просмотр\w*\s+фильм\w*|чтени\w+\s+вслух|обсуждали\s+тему|тренировочн\w+\s+квартир\w*)(?![а-яё])",
        "Кулинарная мастерская, киноклуб, ориентирование в городе, компьютерные и языковые уроки, настольные игры"
    ),
    (
        "Спектакли, концерты и музыкально-театральные занятия",
        r"(?<![а-яё])((?<!маленький\s)спектакл\w*|театральн\w+\s+студи\w*|театральн\w+\s+мастерск\w*|театральн\w+\s+заняти\w*|постановк\w*|концерт\w*|музыкальн\w+\s+вечер\w*|музыкальн\w+\s+заняти\w*|барабан\w*|капустник\w*|цирков\w+\s+представлени\w*|поход\w*\s+в\s+театр|поход\w*\s+в\s+цирк|клоунск\w+\s+выезд\w*)(?![а-яё])",
        r"(?<![а-яё])(репетици\w*|в\s+театр[еу]|гитар\w*|вокал\w*|пели\s+песни|выступлени\w+\s+на\s+сцен\w*)(?![а-яё])",
        "Постановки театральной студии, музыкальные вечера, занятия на барабанах и гитаре, концерты и клоунские выезды"
    ),
    (
        "Праздники, дни рождения и клубные встречи",
        r"(?<![а-яё])(отпразднова\w+|праздновани\w+|рождественск\w+\s+елк\w*|новогодн\w+\s+елк\w*|новогодн\w+\s+праздник\w*|маслениц\w*|масленичн\w+|пасхальн\w+\s+праздник\w*|выпускн\w+\s+вечер\w*|субботн\w+\s+встреч\w*|салон\w*\s+по\s+субботам|клуб\w*\s+выпускников|день\s+рождения\s+студи\w*|день\s+рождения\s+фонд\w*|устроили\s+день\s+рождения|устроили\s+праздник|пикник\w*)(?![а-яё])",
        r"(?<![а-яё])(именинник\w*|чаепити\w*|встреч\w+\s+друзей|вечеринк\w*|дискотек\w*)(?![а-яё])",
        "Сезонные праздники (Масленица, Новый год, Рождество), дни рождения фонда и участников, субботние салоны и пикники"
    ),
    (
        "Спортивные и активные мероприятия",
        r"(?<![а-яё])(футбольн\w+|футбол\w*|забег\w*|марафон\w*|бегущ\w+\s+сердц\w*|красочн\w+\s+забег\w*|велопробег\w*|ката\w+\s+на\s+коньках|на\s+каток|на\s+катке|ката\w+\s+на\s+лыжах|ката\w+\s+на\s+роликах|плавани\w+\s+в\s+бассейн\w*|чемпионат\w*|турнир\w*|батут\w*|скалодром\w*|иппотерапи\w*|конн\w+\s+спорт\w*|танцевальн\w+\s+заняти\w*|танцевальн\w+\s+мастер-класс\w*)(?![а-яё])",
        r"(?<![а-яё])(спортивн\w+\s+тренировк\w*|футбольн\w+\s+тренировк\w*|матч\w*|велосипед\w*|верхов\w+\s+езд\w*|танцевал\w+|йог\w*)(?![а-яё])",
        "Благотворительные забеги, футбольные тренировки и турниры, танцы, каток, конный спорт и активный отдых"
    ),
    (
        "Волонтёрские встречи, обучение и дни открытых дверей",
        r"(?<![а-яё])(школ\w+\s+волонтер\w*|школ\w+\s+волонтёр\w*|набор\w*\s+волонтер\w*|набор\w*\s+волонтёр\w*|встреч\w+\s+волонтер\w*|встреч\w+\s+волонтёр\w*|тренинг\w*\s+для\s+волонтер\w*|супервизи\w*|день\s+открытых\s+дверей|знакомств\w+\s+с\s+фондом|субботник\w*)(?![а-яё])",
        r"(?<![а-яё])(обучени\w+\s+волонтер\w*|обучени\w+\s+волонтёр\w*|семинар\w*\s+для\s+ведущих)(?![а-яё])",
        "Школа волонтёров, дни открытых дверей в студии, вводные встречи, супервизии для ведущих и субботники"
    ),
    (
        "Творческие мастер-классы и художественные студии",
        r"(?<![а-яё])(мастер-класс\w*|арт-лаборатори\w*|арт-терапи\w*|рисова\w+|живопис\w*|керамик\w*|гончар\w*|лепк\w+|лепи\w+|валяни\w+|валял\w+|валяем|оригами|мозаик\w*|рукодели\w*|вышивк\w*|вышива\w+|шить\w*|шьем|шьём|батик\w*|столярн\w*|коллаж\w*|аппликаци\w*|роспис\w+|пленэр\w*|мультстуди\w*|анимаци\w*|поделк\w*)(?![а-яё])",
        r"(?<![а-яё])(творческ\w+\s+заняти\w*|творческ\w+\s+мастерск\w*|заняти\w+\s+в\s+студи\w*)(?![а-яё])",
        "Регулярные занятия и мастер-классы: живопись, керамика, валяние из шерсти, мозаика, шитьё, оригами, столярка, анимация"
    ),
]

STRONG_EVENT_MARKER = re.compile(
    r"(?<![а-яё])(приглаша\w+|ждем\s+вас|ждём\s+вас|приходите|состоится|состоял\w+|прошел|прошёл|прошла|прошли|пройдет|пройдёт|провели|проводим|проводили|занимались|занятие|занятия|на\s+занятиях|встретились|встречаемся|собрались|собрали|ездили|поехали|побывали|отправились|приходили|приехали|в\s+гостях|мастер-класс\w*|ярмарк\w*|базар\w*|выставк\w*|экскурси\w*|спектакл\w*|концерт\w*|лагер\w*|чаепити\w*|тренинг\w*|субботник\w*|забег\w*|турнир\w*|аукцион\w*|день\s+открытых\s+дверей|отпразднова\w+|устроили\s+праздник|устроили\s+день\s+рождения)(?![а-яё])",
    re.IGNORECASE
)

NON_EVENT_EXCLUSION = re.compile(
    r"(мы\s+выиграли\s+грант|точка\s+памяти|светлая\s+память|хранит\w*\s+память|памяти\s+ильи|нет\s+с\s+нами|ушел\s+из\s+жизни|ушёл\s+из\s+жизни|соболезнован|сегодня\s+важный\s+день\s*[-—–]\s*день\s+(?:учителя|защиты\s+детей)|отмечают\s+день\s+семьи|приглашаем\s+вас\s+стать\s+частью\s+нашей\s+большой|запускаем\s+сбор\s+на\s+программу|знакомьтесь\s*[-—–]\s*это\s+|вторую\s+серию\s*[-—–]\s*как\s+оля)",
    re.IGNORECASE
)

AUDIENCE_RULES = [
    (
        "Подростки и взрослые с ментальными особенностями и ОВЗ (ПНИ / Социальные дома)",
        r"(?<![а-яё])(пни|психоневрологическ\w+|социальн\w+\s+дом\w*|филимонк\w*|нагатино-садовник\w*|обручевск\w*|дегунин\w*|овз|инвалид\w*|ментальн\w+|особенност\w+\s+развития|ограниченн\w+\s+возможност\w+|колясочн\w*|на\s+коляск\w+|нарушен\w+\s+зрения|нарушен\w+\s+слуха|аутизм\w*|синдром\w*\s+дауна|интернат\w*\s+№\s*(?:11|12|16|20|30))(?![а-яё])",
        "Жители ПНИ №16, №11, №20, №30, социальных домов «Филимонки», «Нагатино-Садовники», «Обручевский», люди с ОВЗ"
    ),
    (
        "Воспитанники ЦССВ, ДДИ и детских домов-интернатов",
        r"(?<![а-яё])(цссв|дди|детск\w+\s+дом\w*|интернат\w*\s+№\s*(?!11|12|16|20|30)|юнон\w+|спутник\w*|радуг\w+|сколковск\w+|никулин\w+|полярн\w+\s+звезд\w*|алые\s+паруса|каховск\w+\s+ромашк\w*|кунцевск\w*|соколенок|соколёнок|надежд\w+|учм\w+|сирот\w*|оставш\w+\s+без\s+попечения)(?![а-яё])",
        "Дети и подростки с опытом сиротства из ЦССВ «Юнона», «Спутник», «Радуга», «Сколковский», «Никулино», ДДИ №7, интерната №31"
    ),
    (
        "Широкая аудитория, гости и благотворители (открытые события)",
        r"(?<![а-яё])(приглашаем\s+вас|ждем\s+вас|ждём\s+вас|вход\s+свободный|открыт\w+\s+мастер-класс\w*|открыт\w+\s+киноклуб\w*|все\s+желающ\w+|посетител\w*|покупател\w*|гост[иейям]\s+ярмарк\w*|гост[иейям]\s+выставк\w*|гост[иейям]\s+студи\w*|друзья\s+фонда|жител\w+\s+тверского\s+района)(?![а-яё])",
        "Гости благотворительных ярмарок, выставок, открытых мастер-классов и спектаклей, сторонники и доноры фонда"
    ),
    (
        "Выпускники интернатов и детских домов (молодые взрослые)",
        r"(?<![а-яё])(выпускник\w*|клуб\w*\s+выпускников|старш\w+\s+ребят\w*|молод\w+\s+взросл\w*|самостоятельн\w+\s+жизн\w*|тренировочн\w+\s+квартир\w*)(?![а-яё])",
        "Выпускники детских домов и ЦССВ, участники программы сопровождения и проекта «Тренировочная квартира»"
    ),
    (
        "Дети и семьи (в т.ч. приёмные семьи и программа в Беслане)",
        r"(?<![а-яё])(беслан\w*|приемн\w+\s+сем\w+|приёмн\w+\s+сем\w+|кризисн\w+\s+сем\w+|многодетн\w+\s+сем\w*|родител\w+\s+и\s+дет\w+|школьник\w*\s+беслан\w*)(?![а-яё])",
        "Дети и школьники Беслана (выездные смены), приёмные, опекунские и кризисные семьи с детьми"
    ),
    (
        "Волонтёры, тьюторы и команда НКО",
        r"(?<![а-яё])(школ\w+\s+волонтер\w*|школ\w+\s+волонтёр\w*|набор\w*\s+волонтер\w*|набор\w*\s+волонтёр\w*|встреч\w+\s+волонтер\w*|встреч\w+\s+волонтёр\w*|тренинг\w*\s+для\s+волонтер\w*|для\s+ведущих|для\s+кураторов|супервизи\w*)(?![а-яё])",
        "Новые и действующие волонтёры, ведущие творческих мастерских, тьюторы и кураторы программ"
    ),
]

SKILL_RULES = [
    (
        "Художественно-прикладные и ремесленные навыки",
        r"(?<![а-яё])(рисова\w+|живопис\w*|акварел\w*|гуаш\w*|керамик\w*|гончар\w*|лепк\w+|лепи\w+|глин\w+|валяни\w+|валял\w+|шерст\w+|оригами|мозаик\w*|рукодели\w*|вышив\w+|шить\w*|шьем|шьём|ткани|батик\w*|столяр\w*|дерев\w+|коллаж\w*|аппликаци\w*|роспис\w*|поделк\w*|мастери\w+)(?![а-яё])",
        "Живопись и рисунок, керамика и работа с глиной, сухое и мокрое валяние из шерсти, шитьё, мозаика, оригами, столярное дело"
    ),
    (
        "Коммуникативные навыки, эмпатия и тьюторское сопровождение",
        r"(?<![а-яё])(обща\w+|общени\w+|бесед\w+|поддерж\w+|внимани\w+|тепл\w+|друж\w+|инклюзивн\w+|тьютор\w*|индивидуальн\w+|социализаци\w+|адаптаци\w*|психолог\w*|бережно\w*|довери\w*|быть\s+рядом)(?![а-яё])",
        "Бережное общение, эмоциональная поддержка, индивидуальное тьюторство для людей с ОВЗ и опытом сиротства, фасилитация"
    ),
    (
        "Артистические, музыкальные и спортивно-игровые навыки",
        r"(?<![а-яё])(театр\w*|спектакл\w*|репетиц\w*|сцен\w+|музык\w+|концерт\w*|барабан\w*|гитар\w*|песн\w+|пели|танц\w+|игр\w+|викторин\w*|квест\w*|конкурс\w*|спорт\w*|футбол\w*|тренировк\w*|забег\w*|вело\w*|коньк\w*|лыж\w*|клоун\w*)(?![а-яё])",
        "Игра на музыкальных инструментах (гитара, барабаны), актёрское мастерство, проведение игр, викторин, танцев и спорттренировок"
    ),
    (
        "Педагогические и наставнические навыки (кулинария, быт, учёба)",
        r"(?<![а-яё])(учи\w+|обуча\w+|объясня\w+|показыва\w+|преподава\w*|урок\w*|репетитор\w*|английск\w+|компьютер\w*|кулинар\w+|готови\w+|испек\w+|выпека\w+|рецепт\w+|чтени\w+|грамотност\w*|безопасност\w*|ориентировани\w+|ведущ\w+\s+заняти\w*|ведущ\w+\s+мастерск\w*|куратор\w*)(?![а-яё])",
        "Обучение кулинарии, финансовой и бытовой грамотности, ориентированию в городе, английскому языку, работе с ПК и подготовке к школе"
    ),
    (
        "Организационные, ивент- и ярмарочные навыки",
        r"(?<![а-яё])(организ\w+|ярмарк\w*|базар\w*|маркет\w*|стенд\w*|продаж\w*|аукцион\w*|праздник\w*|чаепити\w*|накры\w+\s+стол|встреча\w+\s+гостей|координац\w+|подготовк\w+\s+к|дежур\w*|субботник\w*|уборк\w*)(?![а-яё])",
        "Работа за стендом на благотворительных базарах, презентация сувениров, организация чаепитий, праздников, логистика и субботники"
    ),
    (
        "Мобильность, выездное сопровождение и автоволонтёрство",
        r"(?<![а-яё])(экскурси\w*|сопровожд\w+|прогулк\w*|выезд\w*|поездк\w*|лагер\w*|поход\w*|коляск\w*|колясочн\w*|автоволонтер\w*|автоволонтёр\w*|машин\w*|перевез\w*|транспорт\w*|водител\w*)(?![а-яё])",
        "Сопровождение в выездных лагерях, походах и музеях, физическая помощь маломобильным участникам на колясках, автоволонтёрство"
    ),
    (
        "Медиа, фото/видеосъёмка и культурная медиация",
        r"(?<![а-яё])(фотограф\w*|фотосъёмк\w*|фотосъемк\w*|фотоотчет\w*|фотоотчёт\w*|видео\w*|ролик\w*|монтаж\w*|медиатор\w*|экскурсовод\w*|дизайн\w*|иллюстрац\w*)(?![а-яё])",
        "Репортажная фото- и видеосъёмка мероприятий, проведение медиаторских туров и экскурсий по выставкам, монтаж и дизайн"
    ),
]

VOLUNTEER_EXPLICIT_RE = re.compile(
    r"(?<![а-яё])(волонтер\w*|волонтёр\w*|добровол\w*|тьютор\w*|куратор\w*|наставник\w*|ведущ\w+\s+мастерск\w*|ведущ\w+\s+заняти\w*)(?![а-яё])",
    re.IGNORECASE
)

INSTITUTION_EXTRACT_RE = re.compile(
    r"(?:ПНИ|ЦССВ|ДДИ|интернат\w*|Социальн\w+\s+дом\w*)\s*(?:№\s*\d+|«[^»]+»|\"[^\"]+\")|Беслан|ГЭС[-–]2|Тверская\s*15|Переделкино|Сорочан\w*",
    re.IGNORECASE
)


def clean_vk_text(text: str) -> str:
    """Очищает разметку ссылок ВКонтакте вида [club123|Название] -> Название."""
    text = re.sub(r"\[([^\]|]+)\|([^\]]+)\]", r"\2", str(text))
    text = re.sub(r"[\u2500-\u25ff\ufe0f]", "•", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def analyze_posts(input_xlsx: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Загружает посты, выделяет мероприятия и классифицирует их."""
    df = pd.read_excel(input_xlsx)
    records = []

    for _, row in df.dropna(subset=["text"]).iterrows():
        raw_text = str(row["text"]).strip()
        if len(raw_text) < 25:
            continue
        t_low = raw_text.lower()

        if NON_EVENT_EXCLUSION.search(t_low):
            continue
        if not STRONG_EVENT_MARKER.search(t_low):
            continue

        # Оценка категорий мероприятия
        scores = {}
        for cat_name, hi_pat, med_pat, _ in EVENT_CATEGORIES:
            hi_hits = len(re.findall(hi_pat, t_low))
            med_hits = len(re.findall(med_pat, t_low))
            score = hi_hits * 3 + med_hits * 1
            if score > 0:
                scores[cat_name] = score

        if not scores:
            continue

        sorted_cats = sorted(scores.items(), key=lambda x: x[1], reverse=True)
        primary_type = sorted_cats[0][0]
        all_types = "; ".join([c for c, _ in sorted_cats])

        # Целевая аудитория
        matched_auds = [
            label for label, pat, _ in AUDIENCE_RULES if re.search(pat, t_low)
        ]
        if not matched_auds:
            matched_auds = ["Подопечные фонда «Дети Марии» (общая группа детей и подростков)"]

        # Упомянутые учреждения / площадки
        insts = []
        for m in INSTITUTION_EXTRACT_RE.findall(raw_text):
            norm_m = re.sub(r"\s+", " ", m.strip()).replace('"', '«')
            if norm_m not in insts:
                insts.append(norm_m)

        # Навыки волонтёров
        matched_skills = [
            label for label, pat, _ in SKILL_RULES if re.search(pat, t_low)
        ]
        if not matched_skills:
            matched_skills = ["Коммуникативные навыки, эмпатия и тьюторское сопровождение"]

        explicit_vol = "Да (в тексте)" if VOLUNTEER_EXPLICIT_RE.search(t_low) else "По формату занятия"
        clean_snippet = clean_vk_text(raw_text)[:280]
        if len(raw_text) > 280:
            clean_snippet += "..."

        date_str = str(row["date"])[:10] if pd.notna(row["date"]) else ""
        year_val = int(date_str[:4]) if date_str[:4].isdigit() else None

        records.append({
            "id": int(row["id"]) if pd.notna(row["id"]) else 0,
            "date": date_str,
            "year": year_val,
            "url": str(row["url"]).strip() if pd.notna(row["url"]) else "",
            "primary_type": primary_type,
            "all_types": all_types,
            "audiences_list": matched_auds,
            "audiences": "; ".join(matched_auds),
            "institutions": ", ".join(insts[:5]) if insts else "Студия фонда / выездная площадка",
            "skills_list": matched_skills,
            "skills": "; ".join(matched_skills),
            "explicit_volunteer": explicit_vol,
            "snippet": clean_snippet,
        })

    events_df = pd.DataFrame(records)
    return df, events_df


def style_merged_range(ws, cell_range, font=None, fill=None, alignment=None, border=None):
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if alignment:
                cell.alignment = alignment
            if border:
                cell.border = border


def build_excel_report(df_all: pd.DataFrame, events_df: pd.DataFrame, output_path: str) -> None:
    """Создаёт профессионально оформленный Excel-отчёт с аналитикой и базой мероприятий."""
    wb = openpyxl.Workbook()

    # Цветовая палитра (глубокий сине-зелёный / морской слейт)
    PRIMARY_HEX = "1B4F72"
    SECONDARY_HEX = "2E86C1"
    ACCENT_BG = "EBF5FB"
    ALT_ROW_HEX = "F8F9F9"
    BORDER_HEX = "D5D8DC"

    font_family = "Calibri"
    title_font = Font(name=font_family, size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="D6EAF8")
    sec_hdr_font = Font(name=font_family, size=12, bold=True, color=PRIMARY_HEX)
    tbl_hdr_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    body_font = Font(name=font_family, size=10, bold=False, color="1C2833")
    bold_font = Font(name=font_family, size=10, bold=True, color="1C2833")
    link_font = Font(name=font_family, size=10, underline="single", color="1155CC")

    primary_fill = PatternFill(start_color=PRIMARY_HEX, end_color=PRIMARY_HEX, fill_type="solid")
    secondary_fill = PatternFill(start_color=SECONDARY_HEX, end_color=SECONDARY_HEX, fill_type="solid")
    accent_fill = PatternFill(start_color=ACCENT_BG, end_color=ACCENT_BG, fill_type="solid")
    alt_fill = PatternFill(start_color=ALT_ROW_HEX, end_color=ALT_ROW_HEX, fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_HEX),
        right=Side(style="thin", color=BORDER_HEX),
        top=Side(style="thin", color=BORDER_HEX),
        bottom=Side(style="thin", color=BORDER_HEX),
    )
    total_border = Border(
        top=Side(style="thin", color=PRIMARY_HEX),
        bottom=Side(style="double", color=PRIMARY_HEX),
        left=Side(style="thin", color=BORDER_HEX),
        right=Side(style="thin", color=BORDER_HEX),
    )

    # =========================================================================
    # ЛИСТ 1: СВОДКА И АНАЛИТИКА
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Сводка и аналитика"
    ws1.views.sheetView[0].showGridLines = True
    ws1.page_setup.orientation = "landscape"
    ws1.sheet_properties.pageSetUpPr.fitToPage = True

    # Шапка листа (A1:F2)
    ws1.merge_cells("A1:F1")
    ws1["A1"] = "АНАЛИЗ МЕРОПРИЯТИЙ БЛАГОТВОРИТЕЛЬНОГО ФОНДА «ДЕТИ МАРИИ» (2011–2026)"
    style_merged_range(ws1, "A1:F1", font=title_font, fill=primary_fill, alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:F2")
    ws1["A2"] = "Автоматическая классификация постов ВК: типы мероприятий, целевая аудитория и профиль навыков волонтёров"
    style_merged_range(ws1, "A2:F2", font=subtitle_font, fill=primary_fill, alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws1.row_dimensions[2].height = 20

    total_posts = len(df_all)
    non_empty_posts = int(df_all["text"].notna().sum())
    total_events = len(events_df)
    last_event_row = total_events + 4  # На листе 'База мероприятий' данные начинаются с 5-й строки

    # KPI-карточки (Строки 4–6, 3 карточки по 2 колонки: A:B, C:D, E:F)
    kpis = [
        ("A4:B4", "ВСЕГО ПОСТОВ С ТЕКСТОМ", "A5:B5", non_empty_posts, "#,##0", "A6:B6", f"Всего строк в выгрузке: {total_posts}"),
        ("C4:D4", "ВЫДЕЛЕНО МЕРОПРИЯТИЙ", "C5:D5", f"=COUNTA('База мероприятий'!A5:A{last_event_row})", "#,##0", "C6:D6", "Анонсы, отчёты о занятиях, выездах и ярмарках"),
        ("E4:F4", "ДОЛЯ МЕРОПРИЯТИЙ В КОНТЕНТЕ", "E5:F5", "=C5/A5", "0.0%", "E6:F6", "От всех текстовых публикаций сообщества"),
    ]
    kpi_hdr_font = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    kpi_val_font = Font(name=font_family, size=15, bold=True, color=PRIMARY_HEX)
    kpi_sub_font = Font(name=font_family, size=9, color="566573")
    center_align = Alignment(horizontal="center", vertical="center")

    for h_rng, h_txt, v_rng, v_val, num_fmt, s_rng, s_txt in kpis:
        ws1.merge_cells(h_rng)
        ws1[h_rng.split(":")[0]] = h_txt
        style_merged_range(ws1, h_rng, font=kpi_hdr_font, fill=secondary_fill, alignment=center_align, border=thin_border)

        ws1.merge_cells(v_rng)
        v_cell = ws1[v_rng.split(":")[0]]
        v_cell.value = v_val
        style_merged_range(ws1, v_rng, font=kpi_val_font, fill=accent_fill, alignment=center_align, border=thin_border)
        v_cell.number_format = num_fmt

        ws1.merge_cells(s_rng)
        ws1[s_rng.split(":")[0]] = s_txt
        style_merged_range(ws1, s_rng, font=kpi_sub_font, fill=accent_fill, alignment=center_align, border=thin_border)

    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 26
    ws1.row_dimensions[6].height = 18

    # -------------------------------------------------------------------------
    # ТАБЛИЦА 1: ТИПЫ МЕРОПРИЯТИЙ (Строки 9–20)
    # -------------------------------------------------------------------------
    ws1.merge_cells("A9:F9")
    ws1["A9"] = "1. Распределение мероприятий по типам (основной фокус и все упоминания)"
    ws1["A9"].font = sec_hdr_font
    ws1.row_dimensions[9].height = 22

    t1_headers = [
        "Тип мероприятия",
        "Постов (осн. тип)",
        "Доля (осн. тип)",
        "Всего упоминаний",
        "Характерные форматы в «Дети Марии»",
        "Пример поста ВК"
    ]
    for col_idx, h_text in enumerate(t1_headers, start=1):
        c = ws1.cell(row=10, column=col_idx, value=h_text)
        c.font = tbl_hdr_font
        c.fill = primary_fill
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = thin_border
    ws1.row_dimensions[10].height = 24

    primary_counts = Counter(events_df["primary_type"])
    multi_type_counts = Counter()
    for t_str in events_df["all_types"]:
        for part in t_str.split("; "):
            if part:
                multi_type_counts[part] += 1

    cat_desc_map = {name: desc for name, _, _, desc in EVENT_CATEGORIES}
    sorted_type_rows = sorted(EVENT_CATEGORIES, key=lambda x: primary_counts[x[0]], reverse=True)

    r_idx = 11
    for cat_name, _, _, desc in sorted_type_rows:
        p_cnt = primary_counts[cat_name]
        m_cnt = multi_type_counts[cat_name]
        sample_df = events_df[events_df["primary_type"] == cat_name]
        sample_url = sample_df.iloc[0]["url"] if len(sample_df) > 0 else ""
        sample_id = sample_df.iloc[0]["id"] if len(sample_df) > 0 else ""

        ws1.cell(row=r_idx, column=1, value=cat_name)
        ws1.cell(row=r_idx, column=2, value=p_cnt)
        ws1.cell(row=r_idx, column=3, value=f"=B{r_idx}/$B$20")
        ws1.cell(row=r_idx, column=4, value=m_cnt)
        ws1.cell(row=r_idx, column=5, value=desc)
        link_cell = ws1.cell(row=r_idx, column=6)
        if sample_url:
            link_cell.value = f'=HYPERLINK("{sample_url}", "Пост #{sample_id}")'

        for c_i in range(1, 7):
            c = ws1.cell(row=r_idx, column=c_i)
            c.font = link_font if c_i == 6 else body_font
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = alt_fill
            c.alignment = Alignment(
                horizontal="left" if c_i in (1, 5) else ("center" if c_i == 6 else "right"),
                vertical="center",
                wrap_text=(c_i in (1, 5))
            )
            if c_i in (2, 4):
                c.number_format = "#,##0"
            elif c_i == 3:
                c.number_format = "0.0%"
        ws1.row_dimensions[r_idx].height = 24
        r_idx += 1

    # Итоговая строка Таблицы 1 (Строка 20)
    ws1.cell(row=20, column=1, value="ИТОГО МЕРОПРИЯТИЙ")
    ws1.cell(row=20, column=2, value="=SUM(B11:B19)")
    ws1.cell(row=20, column=3, value="=SUM(C11:C19)")
    ws1.cell(row=20, column=4, value="=SUM(D11:D19)")
    ws1.cell(row=20, column=5, value="Сумма по всем категориям (в 1 посте может быть >1 категории)")
    ws1.cell(row=20, column=6, value="—")
    for c_i in range(1, 7):
        c = ws1.cell(row=20, column=c_i)
        c.font = bold_font
        c.border = total_border
        c.fill = accent_fill
        c.alignment = Alignment(horizontal="left" if c_i in (1, 5) else ("center" if c_i == 6 else "right"), vertical="center")
        if c_i in (2, 4):
            c.number_format = "#,##0"
        elif c_i == 3:
            c.number_format = "0.0%"
    ws1.row_dimensions[20].height = 22

    # -------------------------------------------------------------------------
    # ТАБЛИЦА 2: ЦЕЛЕВАЯ АУДИТОРИЯ МЕРОПРИЯТИЙ (Строки 23–32)
    # -------------------------------------------------------------------------
    ws1.merge_cells("A23:F23")
    ws1["A23"] = "2. Целевая аудитория мероприятий (охват по категориям благополучателей и участников)"
    ws1["A23"].font = sec_hdr_font
    ws1.row_dimensions[23].height = 22

    t2_headers = [
        "Группа целевой аудитории",
        "Кол-во постов",
        "% от мероприятий",
        "Доля от всех упоминаний ЦА",
        "Ключевые упоминаемые учреждения и группы в постах",
        "Пример поста ВК"
    ]
    for col_idx, h_text in enumerate(t2_headers, start=1):
        c = ws1.cell(row=24, column=col_idx, value=h_text)
        c.font = tbl_hdr_font
        c.fill = primary_fill
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = thin_border
    ws1.row_dimensions[24].height = 24

    aud_counts = Counter()
    aud_sample = {}
    for _, erow in events_df.iterrows():
        for a in erow["audiences_list"]:
            aud_counts[a] += 1
            if a not in aud_sample:
                aud_sample[a] = (erow["id"], erow["url"])

    aud_desc_map = {label: desc for label, _, desc in AUDIENCE_RULES}
    aud_desc_map["Подопечные фонда «Дети Марии» (общая группа детей и подростков)"] = (
        "Регулярные участники студийных занятий в Москве без указания конкретного номера интерната в тексте"
    )

    sorted_auds = sorted(aud_counts.items(), key=lambda x: x[1], reverse=True)
    r_idx = 25
    aud_start_row = 25
    aud_end_row = aud_start_row + len(sorted_auds) - 1
    aud_total_row = aud_end_row + 1

    for aud_name, a_cnt in sorted_auds:
        s_id, s_url = aud_sample.get(aud_name, ("", ""))
        ws1.cell(row=r_idx, column=1, value=aud_name)
        ws1.cell(row=r_idx, column=2, value=a_cnt)
        ws1.cell(row=r_idx, column=3, value=f"=B{r_idx}/$B$20")
        ws1.cell(row=r_idx, column=4, value=f"=B{r_idx}/$B${aud_total_row}")
        ws1.cell(row=r_idx, column=5, value=aud_desc_map.get(aud_name, ""))
        link_cell = ws1.cell(row=r_idx, column=6)
        if s_url:
            link_cell.value = f'=HYPERLINK("{s_url}", "Пост #{s_id}")'

        for c_i in range(1, 7):
            c = ws1.cell(row=r_idx, column=c_i)
            c.font = link_font if c_i == 6 else body_font
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = alt_fill
            c.alignment = Alignment(
                horizontal="left" if c_i in (1, 5) else ("center" if c_i == 6 else "right"),
                vertical="center",
                wrap_text=(c_i in (1, 5))
            )
            if c_i == 2:
                c.number_format = "#,##0"
            elif c_i in (3, 4):
                c.number_format = "0.0%"
        ws1.row_dimensions[r_idx].height = 24
        r_idx += 1

    ws1.cell(row=aud_total_row, column=1, value="ВСЕГО УПОМИНАНИЙ АУДИТОРИЙ")
    ws1.cell(row=aud_total_row, column=2, value=f"=SUM(B{aud_start_row}:B{aud_end_row})")
    ws1.cell(row=aud_total_row, column=3, value=f"=B{aud_total_row}/$B$20")
    ws1.cell(row=aud_total_row, column=4, value=f"=SUM(D{aud_start_row}:D{aud_end_row})")
    ws1.cell(row=aud_total_row, column=5, value="Инклюзивные события часто объединяют сразу несколько групп ЦА")
    ws1.cell(row=aud_total_row, column=6, value="—")
    for c_i in range(1, 7):
        c = ws1.cell(row=aud_total_row, column=c_i)
        c.font = bold_font
        c.border = total_border
        c.fill = accent_fill
        c.alignment = Alignment(horizontal="left" if c_i in (1, 5) else ("center" if c_i == 6 else "right"), vertical="center")
        if c_i == 2:
            c.number_format = "#,##0"
        elif c_i in (3, 4):
            c.number_format = "0.0%"
    ws1.row_dimensions[aud_total_row].height = 22

    # -------------------------------------------------------------------------
    # ТАБЛИЦА 3: НАВЫКИ И КОМПЕТЕНЦИИ ВОЛОНТЁРОВ (Строки 35–44)
    # -------------------------------------------------------------------------
    t3_title_row = aud_total_row + 3
    ws1.merge_cells(f"A{t3_title_row}:F{t3_title_row}")
    ws1[f"A{t3_title_row}"] = "3. Востребованные навыки и компетенции волонтёров на мероприятиях"
    ws1[f"A{t3_title_row}"].font = sec_hdr_font
    ws1.row_dimensions[t3_title_row].height = 22

    t3_hdr_row = t3_title_row + 1
    t3_headers = [
        "Категория навыков волонтёров",
        "Кол-во мероприятий",
        "% от мероприятий",
        "Доля в профиле навыков",
        "Конкретные умения и задачи волонтёров",
        "Пример поста ВК"
    ]
    for col_idx, h_text in enumerate(t3_headers, start=1):
        c = ws1.cell(row=t3_hdr_row, column=col_idx, value=h_text)
        c.font = tbl_hdr_font
        c.fill = primary_fill
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = thin_border
    ws1.row_dimensions[t3_hdr_row].height = 24

    skill_counts = Counter()
    skill_sample = {}
    for _, erow in events_df.iterrows():
        for s in erow["skills_list"]:
            skill_counts[s] += 1
            if s not in skill_sample:
                skill_sample[s] = (erow["id"], erow["url"])

    skill_desc_map = {label: desc for label, _, desc in SKILL_RULES}
    sorted_skills = sorted(skill_counts.items(), key=lambda x: x[1], reverse=True)

    r_idx = t3_hdr_row + 1
    sk_start_row = r_idx
    sk_end_row = sk_start_row + len(sorted_skills) - 1
    sk_total_row = sk_end_row + 1

    for sk_name, s_cnt in sorted_skills:
        s_id, s_url = skill_sample.get(sk_name, ("", ""))
        ws1.cell(row=r_idx, column=1, value=sk_name)
        ws1.cell(row=r_idx, column=2, value=s_cnt)
        ws1.cell(row=r_idx, column=3, value=f"=B{r_idx}/$B$20")
        ws1.cell(row=r_idx, column=4, value=f"=B{r_idx}/$B${sk_total_row}")
        ws1.cell(row=r_idx, column=5, value=skill_desc_map.get(sk_name, ""))
        link_cell = ws1.cell(row=r_idx, column=6)
        if s_url:
            link_cell.value = f'=HYPERLINK("{s_url}", "Пост #{s_id}")'

        for c_i in range(1, 7):
            c = ws1.cell(row=r_idx, column=c_i)
            c.font = link_font if c_i == 6 else body_font
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = alt_fill
            c.alignment = Alignment(
                horizontal="left" if c_i in (1, 5) else ("center" if c_i == 6 else "right"),
                vertical="center",
                wrap_text=(c_i in (1, 5))
            )
            if c_i == 2:
                c.number_format = "#,##0"
            elif c_i in (3, 4):
                c.number_format = "0.0%"
        ws1.row_dimensions[r_idx].height = 24
        r_idx += 1

    ws1.cell(row=sk_total_row, column=1, value="ВСЕГО УПОМИНАНИЙ НАВЫКОВ")
    ws1.cell(row=sk_total_row, column=2, value=f"=SUM(B{sk_start_row}:B{sk_end_row})")
    ws1.cell(row=sk_total_row, column=3, value=f"=B{sk_total_row}/$B$20")
    ws1.cell(row=sk_total_row, column=4, value=f"=SUM(D{sk_start_row}:D{sk_end_row})")
    ws1.cell(row=sk_total_row, column=5, value="В каждом мероприятии обычно задействовано от 2 до 4 групп навыков")
    ws1.cell(row=sk_total_row, column=6, value="—")
    for c_i in range(1, 7):
        c = ws1.cell(row=sk_total_row, column=c_i)
        c.font = bold_font
        c.border = total_border
        c.fill = accent_fill
        c.alignment = Alignment(horizontal="left" if c_i in (1, 5) else ("center" if c_i == 6 else "right"), vertical="center")
        if c_i == 2:
            c.number_format = "#,##0"
        elif c_i in (3, 4):
            c.number_format = "0.0%"
    ws1.row_dimensions[sk_total_row].height = 22

    # Встроенная диаграмма под таблицами (Stacked layout)
    chart1 = BarChart()
    chart1.type = "bar"
    chart1.style = 10
    chart1.title = "Распределение мероприятий БФ «Дети Марии» по основным типам"
    chart1.x_axis.title = "Количество постов о мероприятиях"
    chart1.y_axis.title = "Тип мероприятия"
    chart1.width = 22
    chart1.height = 11
    chart1.legend = None

    data_ref = Reference(ws1, min_col=2, min_row=10, max_row=19)
    cats_ref = Reference(ws1, min_col=1, min_row=11, max_row=19)
    chart1.add_data(data_ref, titles_from_data=True)
    chart1.set_categories(cats_ref)
    ws1.add_chart(chart1, f"A{sk_total_row + 3}")

    # Ширина колонок Листа 1
    col_widths_ws1 = {
        "A": 44,
        "B": 18,
        "C": 18,
        "D": 22,
        "E": 52,
        "F": 16,
    }
    for col_l, w in col_widths_ws1.items():
        ws1.column_dimensions[col_l].width = w

    # =========================================================================
    # ЛИСТ 2: БАЗА МЕРОПРИЯТИЙ (ВСЕ ВЫДЕЛЕННЫЕ ПОСТЫ)
    # =========================================================================
    ws2 = wb.create_sheet(title="База мероприятий")
    ws2.views.sheetView[0].showGridLines = True
    ws2.page_setup.orientation = "landscape"
    ws2.sheet_properties.pageSetUpPr.fitToPage = True

    ws2.merge_cells("A1:J1")
    ws2["A1"] = "ПОЛНАЯ БАЗА ВЫДЕЛЕННЫХ МЕРОПРИЯТИЙ ФОНДА «ДЕТИ МАРИИ»"
    style_merged_range(ws2, "A1:J1", font=title_font, fill=primary_fill, alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells("A2:J2")
    ws2["A2"] = f"Отфильтровано {total_events} постов о мероприятиях из {non_empty_posts} текстовых публикаций (2011–2026 гг.)"
    style_merged_range(ws2, "A2:J2", font=subtitle_font, fill=primary_fill, alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws2.row_dimensions[2].height = 20

    ws2_headers = [
        "ID поста",
        "Дата",
        "Ссылка ВК",
        "Основной тип мероприятия",
        "Все категории мероприятия",
        "Целевая аудитория",
        "Упомянутые учреждения / площадки",
        "Востребованные навыки волонтёров",
        "Упоминание волонтёров",
        "Краткий текст поста"
    ]
    for col_idx, h_text in enumerate(ws2_headers, start=1):
        c = ws2.cell(row=4, column=col_idx, value=h_text)
        c.font = tbl_hdr_font
        c.fill = primary_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 3, 9) else "left", vertical="center", wrap_text=True)
        c.border = thin_border
    ws2.row_dimensions[4].height = 24
    ws2.freeze_panes = "A5"

    for idx_offset, (_, erow) in enumerate(events_df.iterrows(), start=5):
        ws2.cell(row=idx_offset, column=1, value=int(erow["id"]))
        ws2.cell(row=idx_offset, column=2, value=erow["date"])
        c_link = ws2.cell(row=idx_offset, column=3)
        if erow["url"]:
            c_link.value = f'=HYPERLINK("{erow["url"]}", "Открыть пост")'
        else:
            c_link.value = "—"
        ws2.cell(row=idx_offset, column=4, value=erow["primary_type"])
        ws2.cell(row=idx_offset, column=5, value=erow["all_types"])
        ws2.cell(row=idx_offset, column=6, value=erow["audiences"])
        ws2.cell(row=idx_offset, column=7, value=erow["institutions"])
        ws2.cell(row=idx_offset, column=8, value=erow["skills"])
        ws2.cell(row=idx_offset, column=9, value=erow["explicit_volunteer"])
        ws2.cell(row=idx_offset, column=10, value=erow["snippet"])

        for c_i in range(1, 11):
            c = ws2.cell(row=idx_offset, column=c_i)
            c.font = link_font if c_i == 3 else body_font
            c.border = thin_border
            if idx_offset % 2 == 0:
                c.fill = alt_fill
            c.alignment = Alignment(
                horizontal="right" if c_i == 1 else ("center" if c_i in (2, 3, 9) else "left"),
                vertical="top",
                wrap_text=(c_i >= 4)
            )
            if c_i == 1:
                c.number_format = "#,##0"
        ws2.row_dimensions[idx_offset].height = 30

    col_widths_ws2 = {
        "A": 12,
        "B": 13,
        "C": 15,
        "D": 34,
        "E": 36,
        "F": 38,
        "G": 28,
        "H": 38,
        "I": 20,
        "J": 50,
    }
    for col_l, w in col_widths_ws2.items():
        ws2.column_dimensions[col_l].width = w

    out_p = pathlib.Path(output_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_p))
    print(f"Отчёт успешно сохранён: {out_p} (выделено мероприятий: {total_events})")


def main():
    parser = argparse.ArgumentParser(description="Анализ мероприятий БФ «Дети Марии» по постам ВК")
    parser.add_argument(
        "--input",
        default=str(REPO_ROOT / "mariaschildren_posts.xlsx"),
        help="Путь к входному Excel-файлу с постами"
    )
    parser.add_argument(
        "--output",
        default=str(REPO_ROOT / "materials" / "02-фонд" / "вк-группа" / "Анализ мероприятий ВК 2011–2026.xlsx"),
        help="Путь к выходному Excel-отчёту"
    )
    args = parser.parse_args()  # <-- парсер разбирает аргументы

    df_all, events_df = analyze_posts(args.input)
    build_excel_report(df_all, events_df, args.output)


if __name__ == "__main__":
    main()