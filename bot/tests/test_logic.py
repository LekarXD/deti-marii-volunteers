"""Проверка логики без Telegram: python -m tests.test_logic (из папки bot/)."""
import asyncio
import os
import sys
import tempfile
from datetime import date, datetime, timedelta
from io import BytesIO

os.environ["DB_PATH"] = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ["BOT_TOKEN"] = "test"
os.environ["QUIET_HOURS"] = "22-9"

from openpyxl import load_workbook  # noqa: E402

from opora_bot import db, exports, jobs, notify, seed, ui  # noqa: E402
from opora_bot.config import TZ  # noqa: E402
from opora_bot.coordinator import candidates  # noqa: E402
from opora_bot.schedule import occurrences, occurs, parse_day  # noqa: E402


class FakeBot:
    def __init__(self):
        self.sent = []

    async def send_message(self, chat_id, text, reply_markup=None, disable_web_page_preview=None):
        self.sent.append((chat_id, text, reply_markup))

    async def send_photo(self, chat_id, photo, caption=None, reply_markup=None):
        self.sent.append((chat_id, f"[photo {photo}] {caption}", reply_markup))


def at(y, m, d, h, mi=0):
    return datetime(y, m, d, h, mi, tzinfo=TZ)


def main():
    db.init()
    seed.seed()
    seed.seed()  # повторный запуск не дублирует
    assert len(db.events()) == len(seed.EVENTS)

    art = db.event("art")
    assert occurs(art, date(2026, 10, 24)) and not occurs(art, date(2026, 10, 23))
    assert not occurs(art, date(2026, 8, 29)), "серия начинается в сентябре"
    occ = occurrences(db.events(), date(2026, 10, 19), date(2026, 10, 26))
    titles = [(e["id"], d.isoformat()) for e, d in occ]
    assert ("excursion", "2026-10-22") in titles and ("nevidimy", "2026-10-26") in titles and ("art", "2026-10-24") in titles
    assert parse_day("31.10", date(2026, 10, 21)) == date(2026, 10, 31)
    assert parse_day("05.01", date(2026, 10, 21)) == date(2027, 1, 5)

    # люди: координатор, допущенная волонтёрка, кандидат
    db.ensure_user(1, "coord"); db.update_user(1, is_coordinator=1, name="Координатор", status="active")
    db.ensure_user(2, "anna"); db.update_user(2, name="Анна", status="active", roles="assistant", days="5")
    db.ensure_user(3, "ivan"); db.update_user(3, name="Иван", status="active", roles="fair", days="1")
    db.ensure_user(4, "new"); db.update_user(4, name="Кандидат", status="candidate")
    day = date(2026, 10, 24)
    pick = [u["id"] for u, why in candidates(art, day)]
    assert pick[0] == 2 and 4 not in pick, pick  # Анна первая: роль и суббота; кандидата не зовём

    # приглашение → согласие → подтверждение
    db.set_invite("art", day.isoformat(), 2, "invited")
    db.set_invite("art", day.isoformat(), 2, "accepted")
    assert db.confirmed_count("art", day.isoformat()) == 0
    db.set_invite("art", day.isoformat(), 2, "confirmed")
    assert db.confirmed_count("art", day.isoformat()) == 1 and ui.coverage(art, day) == "1 из 3"
    assert 2 not in [u["id"] for u, _ in candidates(art, day)], "уже приглашённых не предлагаем"
    assert "Саша" in ui.materials_text(art, day, with_group=True) and "Саша" not in ui.materials_text(art, day, with_group=False)

    bot = FakeBot()

    async def run_at(moment):
        db.now = lambda: moment  # подменяем часы
        await jobs.tick(bot)

    # тихие часы: в 23:00 необязательное сообщение уходит в очередь, утром отправляется
    db.now = lambda: at(2026, 10, 21, 23)
    asyncio.run(notify.send(bot, 2, "ночное", None))
    assert not bot.sent and db.one("SELECT COUNT(*) AS n FROM outbox")["n"] == 1
    asyncio.run(run_at(at(2026, 10, 22, 9, 1)))
    assert any(t == "ночное" for _, t, _ in bot.sent)
    db.now = lambda: at(2026, 10, 21, 23, 30)
    asyncio.run(notify.send(bot, 2, "афиша", None, photo="PIC1"))
    assert db.one("SELECT photo FROM outbox")["photo"] == "PIC1"
    asyncio.run(run_at(at(2026, 10, 22, 9, 2)))
    assert any(t == "[photo PIC1] афиша" for _, t, _ in bot.sent), "картинка ночью ждёт до утра"

    # напоминание накануне в 18:00, утром в день, вопрос после занятия
    bot.sent.clear()
    asyncio.run(run_at(at(2026, 10, 23, 17)))
    assert not [t for c, t, _ in bot.sent if c == 2 and "Завтра" in t]
    asyncio.run(run_at(at(2026, 10, 23, 18, 5)))
    asyncio.run(run_at(at(2026, 10, 23, 19)))  # повтор не дублирует
    assert len([t for c, t, _ in bot.sent if c == 2 and "Завтра" in t]) == 1
    asyncio.run(run_at(at(2026, 10, 24, 9, 30)))
    assert not [t for c, t, _ in bot.sent if c == 2 and "Сегодня" in t], "утреннего повтора больше нет"
    asyncio.run(run_at(at(2026, 10, 24, 19, 5)))
    assert [t for c, t, _ in bot.sent if c == 2 and "Как прошло" in t]

    # координатору: молчание 48 часов — в ленту, не сообщением; утром одна сводка с нехваткой людей
    bot.sent.clear()
    db.set_invite("art", "2026-10-31", 3, "invited")
    db.run("UPDATE invites SET invited_at=? WHERE user_id=3", at(2026, 10, 22, 10).isoformat())
    db.run("DELETE FROM sent WHERE key LIKE 'digest:%'")
    asyncio.run(run_at(at(2026, 10, 25, 10, 5)))
    asyncio.run(run_at(at(2026, 10, 25, 12, 0)))
    to_coord = [t for c, t, _ in bot.sent if c == 1]
    assert len(to_coord) == 1 and "Сводка" in to_coord[0] and "31.10 Арт-суббота — 0 из 3" in to_coord[0], to_coord
    assert db.one("SELECT 1 FROM activity WHERE kind='noanswer' AND user_id=3"), "молчание — в ленте"

    # собеседование: напоминание за сутки и за 15 минут, без повторов
    bot.sent.clear()
    mid = db.run("INSERT INTO meetings (user_id, starts_at, place) VALUES (?,?,?)", 4, at(2026, 10, 27, 18).isoformat(), "Студия")
    for when in (at(2026, 10, 26, 18, 30), at(2026, 10, 27, 16, 10), at(2026, 10, 27, 17, 46), at(2026, 10, 27, 17, 50)):
        asyncio.run(run_at(when))
    texts = [t for c, t, _ in bot.sent if c == 4]
    assert len(texts) == 2 and "Завтра в 18:00" in texts[0] and "Через 15 минут" in texts[1], texts

    # выгрузки открываются, .ics корректный
    db.set_invite("art", day.isoformat(), 2, "attended", hours=4)
    db.now = lambda: at(2026, 10, 28, 12)
    for blob in (exports.volunteers_xlsx(), exports.teams_xlsx(), exports.hours_xlsx(date(2026, 10, 28)), exports.team_xlsx(art, day)):
        load_workbook(BytesIO(blob))
    hours = load_workbook(BytesIO(exports.hours_xlsx(date(2026, 10, 28)))).active
    assert hours.max_row == 2 and hours["E2"].value == 4
    summary = exports.month_summary(date(2026, 10, 28))
    assert "Вышли волонтёров: 1, часов: 4" in summary, summary
    cal = exports.ics([(art, day)]).decode()
    assert "DTSTART;TZID=Europe/Moscow:20261024T140000" in cal and cal.count("BEGIN:VEVENT") == 1

    # имя из Telegram «=HYPERLINK(…)» в Excel остаётся текстом, а не формулой
    trap = '=HYPERLINK("http://x","y")'
    db.ensure_user(5, "anya"); db.update_user(5, name=trap, status="active")
    cell = next(c for row in load_workbook(BytesIO(exports.volunteers_xlsx())).active.iter_rows() for c in row if c.value == trap)
    assert cell.data_type == "s", cell.data_type
    # «<» в имени не ломает итоги месяца (сообщение в HTML)
    db.update_user(5, name="Аня <3")
    db.set_invite("art", "2026-10-17", 5, "invited")
    db.set_invite("art", "2026-10-17", 5, "declined", answered_at=at(2026, 10, 17, 9).isoformat())
    summary = exports.month_summary(date(2026, 10, 28))
    assert "• Аня &lt;3 — Арт-суббота, 17 октября" in summary and "<3" not in summary, summary

    # тихие часы без перехода через полночь (QUIET_HOURS=0-8): днём сообщения уходят
    notify.QUIET_FROM, notify.QUIET_TO = 0, 8
    assert notify.is_quiet(at(2026, 10, 22, 3)) and not notify.is_quiet(at(2026, 10, 22, 12))
    notify.QUIET_FROM, notify.QUIET_TO = 22, 9
    assert notify.is_quiet(at(2026, 10, 22, 23)) and not notify.is_quiet(at(2026, 10, 22, 12))

    # событие с временем «24:00» (записанное до проверки ввода) не мешает напоминанию о следующем событии
    for eid, start, end in (("bad", "20:00", "24:00"), ("good", "21:00", "22:00")):
        db.run("INSERT INTO events (id, title, start_date, time, end_time, need) VALUES (?,?,?,?,?,1)", eid, eid, "2026-11-18", start, end)
        db.set_invite(eid, "2026-11-18", 2, "confirmed")
    bot.sent.clear()
    jobs.log.disabled = True  # ожидаемая ошибка по «bad» — в журнал, не в вывод теста
    asyncio.run(jobs.event_reminders(bot, at(2026, 11, 17, 18, 5)))
    jobs.log.disabled = False
    assert [t for c, t, _ in bot.sent if c == 2 and "Завтра" in t and "good" in t], bot.sent

    print("PASS: даты и серии, подбор людей, приглашения, материалы и группа, тихие часы, одно напоминание накануне, "
          "лента и утренняя сводка координатору, собеседования (за сутки и за 15 минут), выгрузки Excel и .ics")


if __name__ == "__main__":
    sys.exit(main())
