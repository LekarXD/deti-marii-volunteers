"""Real CSV, PDF and DOCX output from one immutable confirmed report snapshot."""
import csv,io
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.core.pdf_font import font_path

MIME={'csv':'text/csv','pdf':'application/pdf','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
def metric_value(metric):
    if metric['value'] is None:return 'Нет данных'
    if metric['key']=='volunteer_minutes':
        minutes=int(metric['value']);return f'{minutes} мин ({minutes//60} ч {minutes%60:02d} мин)'
    value=metric['value'];value=int(value) if float(value).is_integer() else value
    return f"{value} {metric['unit']}"
def period(snapshot):return f"{snapshot['period_from']} — {snapshot['period_to']}"
def safe_csv(value):
    text=str(value) if value is not None else ''
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) and not isinstance(value,(int,float)) else text

def csv_report(snapshot):
    output=io.StringIO(newline='');writer=csv.writer(output,delimiter=';',lineterminator='\r\n')
    def row(values):writer.writerow([safe_csv(value) for value in values])
    row(['Отчёт',snapshot['title']]);row(['Период',snapshot['period_from'],snapshot['period_to']]);row(['Область',snapshot['metrics']['scope']['label']]);row(['Определения',snapshot['definitions_version']]);row([])
    row(['Показатель','Значение','Единица','Статус','Определение','Ограничение'])
    for metric in snapshot['metrics']['metrics']:row([metric['label'],metric['value'],metric['unit'],metric['status'],metric['definition'],metric['reason']])
    row([]);row(['Дата помощи','Волонтёр','Вид помощи','Разница минут','Источник','Версия источника','Факт','Группа'])
    for fact in snapshot['metrics']['rows']:row([fact['service_date'],fact['display_name'],fact['title'],fact['delta_minutes'],fact['source_kind'],fact['source_revision'],fact['id'],fact['group_title']])
    row([]);row(['Раздел','Текст'])
    for section in snapshot['sections']:row([section['title'],section['text']])
    return b'\xef\xbb\xbf'+output.getvalue().encode('utf-8')

def pdf_report(snapshot):
    if 'OporaReport' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('OporaReport',str(font_path())))
    body=ParagraphStyle('body',fontName='OporaReport',fontSize=10,leading=16,spaceAfter=10)
    heading=ParagraphStyle('heading',fontName='OporaReport',fontSize=18,leading=25,spaceAfter=12,spaceBefore=20)
    title=ParagraphStyle('title',fontName='OporaReport',fontSize=28,leading=36,spaceAfter=24)
    small=ParagraphStyle('small',fontName='OporaReport',fontSize=8,leading=12,textColor=colors.HexColor('#646464'),spaceAfter=8)
    def p(text,style=body):return Paragraph(escape(str(text)).replace('\n','<br/>'),style)
    story=[p(snapshot['title'],title),p('Период '+period(snapshot)),p('Область '+snapshot['metrics']['scope']['label']),p('Определения '+snapshot['definitions_version']+' · снимок '+snapshot['metrics']['generated_at'],small),p('Показатели',heading)]
    table=[[p('Показатель'),p('Значение')]]+[[p(metric['label']),p(metric_value(metric))] for metric in snapshot['metrics']['metrics']]
    block=Table(table,colWidths=[315,168],repeatRows=1,hAlign='LEFT');block.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#FFEB7A')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),12),('RIGHTPADDING',(0,0),(-1,-1),12),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),4),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#deded8'))]));story.append(block)
    story.append(p('Как считаем',heading))
    for metric in snapshot['metrics']['metrics']:
        story.extend([p(metric['label']),p(metric['definition']+(' '+metric['reason'] if metric['reason'] else ''),small)])
    for section in snapshot['sections']:story.extend([p(section['title'],heading),p(section['text'])])
    story.extend([Spacer(1,18),p(f"Отчёт {snapshot['id']} · версия {snapshot['version']}",small),p('Сформирован по подтверждённому снимку данных. Последующие исправления отражаются в новых версиях отчёта.',small)])
    output=io.BytesIO()
    def footer(canvas,doc):canvas.setFont('OporaReport',8);canvas.drawString(56,30,f'Опора · {doc.page}')
    SimpleDocTemplate(output,pagesize=A4,leftMargin=56,rightMargin=56,topMargin=45,bottomMargin=55,title=snapshot['title'],author='Опора').build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()

def docx_report(snapshot):
    document=Document();section=document.sections[0];section.top_margin=Inches(.8);section.bottom_margin=Inches(1.0);section.footer_distance=Inches(.35);section.left_margin=Inches(.85);section.right_margin=Inches(.85)
    normal=document.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(8)
    for name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        style=document.styles[name];style.font.name='DejaVu Sans';style.font.color.rgb=RGBColor(0,0,0)
        ppr=style.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
        rpr=style.element.find(qn('w:rPr'))
        fonts=rpr.find(qn('w:rFonts')) if rpr is not None else None
        if fonts is not None:
            for attribute in list(fonts.attrib):
                if attribute.endswith('Theme'):del fonts.attrib[attribute]
    document.add_paragraph(snapshot['title'],style='Title')
    document.add_paragraph('Период '+period(snapshot));document.add_paragraph('Область '+snapshot['metrics']['scope']['label'])
    document.add_paragraph('Отчёт содержит подтверждённые показатели и проверенные разделы за выбранный период. Ожидающие подтверждения часы в итог не входят.')
    document.add_heading('Показатели',level=1)
    table=document.add_table(rows=1,cols=2);table.autofit=False;table.columns[0].width=Inches(4.0);table.columns[1].width=Inches(2.3)
    table.rows[0].cells[0].text='Показатель';table.rows[0].cells[1].text='Значение'
    for cell in table.rows[0].cells:
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFF4CC');cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:run.bold=True
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for metric in snapshot['metrics']['metrics']:
        cells=table.add_row().cells;cells[0].text=metric['label'];cells[1].text=metric_value(metric)
    for row in table.rows:
        for cell in row.cells:
            margin=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'110');node.set(qn('w:type'),'dxa');margin.append(node)
            cell._tc.get_or_add_tcPr().append(margin)
    document.add_heading('Как считаем',level=1)
    for metric in snapshot['metrics']['metrics']:
        document.add_paragraph(metric['label'],style='Heading 2');document.add_paragraph(metric['definition']+(' '+metric['reason'] if metric['reason'] else ''))
    for content in snapshot['sections']:document.add_heading(content['title'],level=1);document.add_paragraph(content['text'])
    document.add_paragraph(f"Определения {snapshot['definitions_version']} · снимок {snapshot['metrics']['generated_at']}")
    document.add_paragraph(f"Отчёт {snapshot['id']} · версия {snapshot['version']}")
    for cell in document.sections:
        footer=cell.footer.paragraphs[0];footer.text='Опора';footer.style=document.styles['Normal']
        for run in footer.runs:run.font.size=Pt(8)
    output=io.BytesIO();document.save(output);return output.getvalue()

def render(snapshot,format):return {'csv':csv_report,'pdf':pdf_report,'docx':docx_report}[format](snapshot)
