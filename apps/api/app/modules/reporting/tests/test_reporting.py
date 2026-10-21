import asyncio,csv,io,os
from datetime import date,timedelta
from pathlib import Path
from uuid import UUID,uuid4
import pytest,pytest_asyncio
from fastapi import FastAPI
from sqlalchemy import select,func,text
from sqlalchemy.ext.asyncio import create_async_engine,async_sessionmaker
from docx import Document
from pypdf import PdfReader
from app.core.db import Base,get_db
from app.core.config import get_settings
from app.core.clock import utcnow
from app.core.errors import install_error_handlers
from app.core.models import Foundation,User,StaffMembership,PermissionGrant,Session,MediaAsset
from app.core.crypto import digest
from app.core.model_loader import load_models
from app.modules.auth.router import router as auth_router
from app.modules.assets.router import router as asset_router
from app.modules.profile.router import router as profile_router
from app.modules.profile.tests.test_profile import attendance_fixture
from app.modules.profile.service import append_service_time
from app.modules.applications.tests.test_applications import staff,actor,post
from app.modules.activities.assignment_models import Assignment,AttendanceEvent
from app.modules.activities.calendar_models import EventOccurrence
from app.modules.participants.models import ParticipantGroup
from app.modules.reporting.models import ReportDraft,ReportExport
from app.modules.reporting.router import router
load_models()

@pytest_asyncio.fixture
async def env():
    url=os.environ.get('OPORA_REPORTING_TEST_DATABASE_URL','')
    if not url:pytest.skip('Set dedicated OPORA_REPORTING_TEST_DATABASE_URL')
    if not url.endswith('/opora_reporting_test'):raise RuntimeError('Dedicated reporting database required')
    engine=create_async_engine(url);factory=async_sessionmaker(engine,expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.execute(text('CREATE EXTENSION IF NOT EXISTS btree_gist'))
        await conn.run_sync(Base.metadata.drop_all);await conn.run_sync(Base.metadata.create_all)
    async with factory() as db:db.add(Foundation(id=get_settings().foundation_id,name='Фонд «Дети Марии»'));await db.commit()
    app=FastAPI();install_error_handlers(app)
    for routes in (auth_router,asset_router,profile_router,router):app.include_router(routes)
    async def override():
        async with factory() as db:
            try:yield db
            except BaseException:await db.rollback();raise
    app.dependency_overrides[get_db]=override
    yield app,factory
    await engine.dispose()
async def principal(client,factory):
    from app.core.auth import resolve_principal
    async with factory() as db:return await __import__('app.core.auth',fromlist=['resolve_principal']).resolve_principal(db,await db.scalar(select(Session).where(Session.token_hash==digest(client.cookies.get('opora_session')))))
async def reporter(app,factory):return await staff(app,factory,('reports.read','reports.write','reports.export','hours.review','recognition.manage'))
async def group_reporter(app,factory,group_id):
    client=await staff(app,factory,());p=await principal(client,factory)
    async with factory() as db:
        membership=await db.scalar(select(StaffMembership).where(StaffMembership.user_id==p.user_id))
        for name in ('reports.read','reports.write','reports.export'):db.add(PermissionGrant(membership_id=membership.id,permission=name,scope_type='group',scope_id=group_id))
        await db.commit()
    return client
async def fixture_facts(app,factory):
    first=await staff(app,factory,());second=await staff(app,factory,());coord=await reporter(app,factory)
    p1=await principal(first,factory);p2=await principal(second,factory)
    async with factory() as db:
        g1=ParticipantGroup(foundation_id=p1.foundation_id,title='Мастерская А');g2=ParticipantGroup(foundation_id=p1.foundation_id,title='Мастерская Б');db.add_all([g1,g2]);user=await db.get(User,p2.user_id);user.display_name='=FORMULA';await db.commit()
    period_from=date.today()-timedelta(days=20)
    events=[]
    for p,day,group in [(p1,date.today()-timedelta(days=110),g1),(p1,date.today()-timedelta(days=5),g1),(p2,date.today()-timedelta(days=4),g2)]:
        assignment_id,attendance_id,event_id=await attendance_fixture(factory,p.user_id,day)
        async with factory() as db:
            assignment=await db.get(Assignment,assignment_id);assignment.group_id_snapshot=group.id;occurrence=await db.get(EventOccurrence,assignment.occurrence_id);occurrence.group_id=group.id
            await append_service_time(db,p,p.user_id,assignment_id,event_id);await db.commit()
        events.append((p,assignment_id,attendance_id,event_id))
    p,assignment_id,attendance_id,_=events[1]
    async with factory() as db:
        correction=AttendanceEvent(foundation_id=p.foundation_id,assignment_id=assignment_id,attendance_id=attendance_id,user_id=p.user_id,version=2,old_minutes=120,new_minutes=60,reason='Уточнён факт помощи',actor_id='fixture');db.add(correction);await db.flush();await append_service_time(db,p,p.user_id,assignment_id,correction.id);await db.commit()
    manual=(await post(second,'/me/hours/requests',dict(title='Подготовка материалов',service_date=(date.today()-timedelta(days=3)).isoformat(),requested_minutes=45))).json()
    reviewed=await post(coord,f"/coord/hours/requests/{manual['id']}/review",dict(expected_version=1,decision='approved',approved_minutes=45,reason='Подтверждена подготовка материалов'));assert reviewed.status_code==200,reviewed.text
    return dict(first=first,second=second,coord=coord,p1=p1,p2=p2,g1=g1,g2=g2,period_from=period_from,manual=manual)
async def create(client,period_from,group=None):
    response=await post(client,'/coord/reports',dict(scope_type='group' if group else 'foundation',scope_id=str(group) if group else None,period_from=period_from.isoformat(),period_to=date.today().isoformat(),kind='grant',title='Отчёт о волонтёрской помощи'));assert response.status_code==200,response.text
    return response.json()
async def confirmed(client,report):
    for section in report['sections']:
        if section['key']!='stories':
            response=await client.put(f"/api/v1/coord/reports/{report['id']}/sections/{section['key']}",json=dict(expected_version=report['version'],text='Ручной проверенный текст раздела.'),headers={'Idempotency-Key':str(uuid4())});assert response.status_code==200,response.text;report=response.json()
        response=await post(client,f"/coord/reports/{report['id']}/sections/{section['key']}/decision",dict(expected_version=report['version'],decision='confirm'));assert response.status_code==200,response.text;report=response.json()
    response=await post(client,f"/coord/reports/{report['id']}/confirm",dict(expected_version=report['version']));assert response.status_code==200,response.text
    return response.json()

@pytest.mark.asyncio
async def test_canonical_metrics_match_signed_facts_and_actual_scope(env):
    app,factory=env;f=await fixture_facts(app,factory);group=await group_reporter(app,factory,f['g1'].id)
    try:
        all_report=await create(f['coord'],f['period_from']);values={m['key']:m for m in all_report['metrics']['metrics']}
        assert values['volunteer_minutes']['value']==225 and values['held_occurrences']['value']==2 and values['active_volunteers']['value']==2
        assert values['children_attended']['value'] is None and values['children_attended']['status']=='unavailable'
        scoped=await create(group,f['period_from'],f['g1'].id);metrics={m['key']:m for m in scoped['metrics']['metrics']}
        assert metrics['volunteer_minutes']['value']==60 and metrics['held_occurrences']['value']==1 and metrics['retention_90']['value']==100
        assert {row['user_id'] for row in scoped['metrics']['rows']}=={str(f['p1'].user_id)}
        assert (await group.get('/api/v1/coord/reports/'+all_report['id'])).status_code==403
        assert (await group.get('/api/v1/coord/reports')).json()[0]['id']==scoped['id']
        scopes=(await group.get('/api/v1/coord/reports/scopes')).json();assert len(scopes)==1 and scopes[0]['scope_id']==str(f['g1'].id)
        assert 'phone' not in str(scoped['metrics']) and 'about' not in str(scoped['metrics'])
    finally:
        for client in (f['first'],f['second'],f['coord'],group):await client.aclose()

@pytest.mark.asyncio
async def test_reviewed_exports_are_real_and_revocation_blocks_file_and_receipt(env):
    app,factory=env;f=await fixture_facts(app,factory)
    try:
        report=await create(f['coord'],f['period_from']);assert (await post(f['coord'],f"/coord/reports/{report['id']}/exports",dict(expected_version=1,format='csv'))).status_code==409
        report=await confirmed(f['coord'],report);evidence=Path('/private/tmp/opora-reporting-evidence');evidence.mkdir(exist_ok=True)
        csv_export=None
        for format in ('csv','pdf','docx'):
            key='export-'+format;payload=dict(expected_version=report['version'],format=format)
            response=await post(f['coord'],f"/coord/reports/{report['id']}/exports",payload,key);assert response.status_code==200,response.text
            value=response.json();assert value['state']=='ready' and value['asset_id']
            assert (await post(f['coord'],f"/coord/reports/{report['id']}/exports",payload,key)).json()==value
            content=await f['coord'].get('/api/v1/assets/'+value['asset_id']+'/content');assert content.status_code==200,content.text
            (evidence/('report.'+format)).write_bytes(content.content)
            if format=='csv':assert "'=FORMULA" in content.content.decode('utf-8-sig');csv_export=value
            elif format=='pdf':
                text=' '.join(page.extract_text() for page in PdfReader(io.BytesIO(content.content)).pages);assert '225 мин' in text and report['period_from'] in text
            else:
                doc=Document(io.BytesIO(content.content));assert any(p.text=='Отчёт о волонтёрской помощи' for p in doc.paragraphs);assert any('225 мин' in cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        preview=(await f['coord'].get(f"/api/v1/coord/reports/{report['id']}/dobro-preview")).json();assert preview['state']=='unavailable' and not preview['upload_available'] and preview['total_minutes']==225
        async with factory() as db:
            assert await db.scalar(select(func.count()).select_from(ReportExport))==3
            grants=(await db.scalars(select(PermissionGrant).where(PermissionGrant.permission=='reports.export'))).all()
            for grant in grants:grant.revoked_at=utcnow()
            await db.commit()
        assert (await f['coord'].get('/api/v1/assets/'+csv_export['asset_id']+'/content')).status_code==403
        assert (await post(f['coord'],f"/coord/reports/{report['id']}/exports",dict(expected_version=report['version'],format='csv'),'export-csv')).status_code==403
        assert (await f['coord'].get('/api/v1/coord/reports/'+report['id'])).status_code==200
    finally:
        for client in (f['first'],f['second'],f['coord']):await client.aclose()

@pytest.mark.asyncio
async def test_stale_report_facts_force_explicit_refresh_before_confirmation(env):
    app,factory=env;f=await fixture_facts(app,factory)
    try:
        report=await create(f['coord'],f['period_from'])
        updated=await post(f['coord'],f"/coord/hours/requests/{f['manual']['id']}/review",dict(expected_version=2,decision='approved',approved_minutes=30,reason='Уточнено время подготовки'));assert updated.status_code==200,updated.text
        for section in report['sections']:
            if section['key']!='stories':
                response=await f['coord'].put(f"/api/v1/coord/reports/{report['id']}/sections/{section['key']}",json=dict(expected_version=report['version'],text='Проверенный текст'),headers={'Idempotency-Key':str(uuid4())});report=response.json()
            response=await post(f['coord'],f"/coord/reports/{report['id']}/sections/{section['key']}/decision",dict(expected_version=report['version'],decision='confirm'));report=response.json()
        blocked=await post(f['coord'],f"/coord/reports/{report['id']}/confirm",dict(expected_version=report['version']));assert blocked.status_code==409 and blocked.json()['code']=='report_facts_changed'
        refreshed=await post(f['coord'],f"/coord/reports/{report['id']}/refresh",dict(expected_version=report['version']));assert refreshed.status_code==200,refreshed.text
        report=refreshed.json();assert not any(section['confirmed'] for section in report['sections'])
        assert next(m for m in report['metrics']['metrics'] if m['key']=='volunteer_minutes')['value']==210
    finally:
        for client in (f['first'],f['second'],f['coord']):await client.aclose()

@pytest.mark.asyncio
async def test_staff_document_request_uses_canonical_profile_basis_and_rechecks_it(env):
    app,factory=env;f=await fixture_facts(app,factory)
    try:
        bases=(await f['coord'].get('/api/v1/coord/profile/recognition-bases')).json();basis=next(row for row in bases if row['user_id']==str(f['p2'].user_id) and row['title']=='Подготовка материалов')
        payload=dict(kind='letter',user_id=basis['user_id'],basis_entry_id=basis['entry_id'],period_from=f['period_from'].isoformat(),period_to=date.today().isoformat(),recipient='Для университета')
        response=await post(f['coord'],'/coord/profile/documents',payload,'letter');assert response.status_code==200,response.text
        letter=response.json();assert letter['state']=='requested'
        assert (await post(f['coord'],'/coord/profile/documents',payload,'letter')).json()['id']==letter['id']
        assert (await f['second'].get('/api/v1/me/profile/documents')).json()[0]['id']==letter['id']
        corrected=await post(f['coord'],f"/coord/hours/requests/{f['manual']['id']}/review",dict(expected_version=2,decision='rejected',approved_minutes=0,reason='Отменено ошибочное подтверждение'));assert corrected.status_code==200
        blocked=await post(f['coord'],f"/coord/profile/documents/{letter['id']}/review",dict(expected_version=1,decision='approved',body='Спасибо за помощь',public_response='Письмо проверено'));assert blocked.status_code==409 and blocked.json()['code']=='document_basis_invalid'
    finally:
        for client in (f['first'],f['second'],f['coord']):await client.aclose()
